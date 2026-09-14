#!/usr/bin/env python3
"""Genera la presentazione di vendita di "Un campo alla volta"."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image
import os
import subprocess

HERE   = os.path.dirname(os.path.abspath(__file__))
SHOTS  = os.path.join(HERE, "schermate")
BUILD  = os.path.join(HERE, ".build")          # ritagli intermedi, non versionati
OUT    = os.path.join(HERE, "un-campo-alla-volta.pptx")

BLU      = RGBColor(0x00, 0x4A, 0x8F)
BLU_CH   = RGBColor(0x2E, 0x7C, 0xC4)
INK      = RGBColor(0x14, 0x14, 0x16)
GRIGIO   = RGBColor(0x5A, 0x5F, 0x66)
CARTA    = RGBColor(0xF7, 0xF7, 0xF5)
BIANCO   = RGBColor(0xFF, 0xFF, 0xFF)
AMBRA    = RGBColor(0xB4, 0x53, 0x09)
AMBRA_BG = RGBColor(0xFD, 0xF6, 0xE3)
VERDE    = RGBColor(0x1B, 0x7F, 0x4B)

def prepara_ritagli():
    """Deriva i ritagli usati nelle slide dalle schermate in schermate/.

    Le schermate sono catturate dall'app in esecuzione (vedi README).
    I ritagli sono intermedi: si rigenerano, non si versionano.
    """
    os.makedirs(BUILD, exist_ok=True)

    # Il PDF compilato: ne ricavo un PNG e ne tengo la fascia leggibile.
    pdf = os.path.join(SHOTS, "04-modulo-compilato.pdf")
    png = os.path.join(BUILD, "modulo-compilato.png")
    if not os.path.exists(png):
        subprocess.run(["qlmanage", "-t", "-s", "1600", "-o", BUILD, pdf],
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, check=False)
        prodotto = os.path.join(BUILD, "04-modulo-compilato.pdf.png")
        if os.path.exists(prodotto):
            os.replace(prodotto, png)
    if os.path.exists(png):
        im = Image.open(png); w, h = im.size
        im.crop((int(w * .05), int(h * .06), int(w * .80), int(h * .38))) \
          .save(os.path.join(BUILD, "pdf-compilato-crop.png"))

    # La schermata di upload: tolgo le fasce vuote sopra e sotto.
    up = Image.open(os.path.join(SHOTS, "01-upload.png")); W_, H_ = up.size
    up.crop((0, int(H_ * .22), W_, int(H_ * .80))) \
      .save(os.path.join(BUILD, "upload-crop.png"))


W, H = Inches(13.333), Inches(7.5)
prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]

prepara_ritagli()


def slide(bg=CARTA):
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    r.fill.solid(); r.fill.fore_color.rgb = bg
    r.line.fill.background(); r.shadow.inherit = False
    return s


def text(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         spacing=1.0, space_after=0):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    first = True
    for item in runs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.line_spacing = item.get("ls", spacing)
        p.space_after = Pt(item.get("sa", space_after))
        p.space_before = Pt(item.get("sb", 0))
        r = p.add_run()
        r.text = item["t"]
        f = r.font
        f.name = item.get("font", "Helvetica Neue")
        f.size = Pt(item["sz"])
        f.bold = item.get("b", False)
        f.italic = item.get("i", False)
        f.color.rgb = item.get("c", INK)
    return tb


def pic(s, path, x, y, max_w, max_h, border=True):
    """Inserisce l'immagine scalata dentro il box, centrata."""
    iw, ih = Image.open(path).size
    k = min(max_w / iw, max_h / ih)
    w, h = int(iw * k), int(ih * k)
    px = x + int((max_w - w) / 2)
    py = y + int((max_h - h) / 2)
    if border:
        bd = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, px - Emu(9525), py - Emu(9525),
                                w + Emu(19050), h + Emu(19050))
        bd.fill.solid(); bd.fill.fore_color.rgb = BIANCO
        bd.line.color.rgb = RGBColor(0xD8, 0xD8, 0xD4); bd.line.width = Pt(0.75)
        bd.shadow.inherit = False
    s.shapes.add_picture(path, px, py, w, h)


def box(s, x, y, w, h, fill, line=None, radius=True):
    shp = s.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line:
        shp.line.color.rgb = line; shp.line.width = Pt(1)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    if radius:
        try:
            shp.adjustments[0] = 0.06
        except Exception:
            pass
    return shp


def rule(s, x, y, w, color=BLU, thick=3):
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, Pt(thick))
    r.fill.solid(); r.fill.fore_color.rgb = color
    r.line.fill.background(); r.shadow.inherit = False


def kicker(s, x, y, testo, color=BLU):
    text(s, x, y, Inches(9), Inches(0.3),
         [{"t": testo.upper(), "sz": 12, "b": True, "c": color}])


def notes(s, txt):
    s.notes_slide.notes_text_frame.text = txt


# ─────────────────────────────────────────────── 1. Titolo
s = slide(BIANCO)
band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.28), H)
band.fill.solid(); band.fill.fore_color.rgb = BLU
band.line.fill.background(); band.shadow.inherit = False

text(s, Inches(1.1), Inches(2.35), Inches(10.5), Inches(1.4),
     [{"t": "Un campo alla volta", "sz": 60, "b": True, "c": INK}])
text(s, Inches(1.1), Inches(3.6), Inches(9.6), Inches(1.2),
     [{"t": "Chi non capisce il linguaggio burocratico compila il modulo da solo, "
            "fino in fondo.", "sz": 24, "c": GRIGIO, "ls": 1.25}])
rule(s, Inches(1.1), Inches(5.05), Inches(1.6))
text(s, Inches(1.1), Inches(5.45), Inches(10), Inches(0.6),
     [{"t": "Carica il PDF  ·  Rispondi a una domanda alla volta  ·  Scarica il modulo compilato",
       "sz": 15, "b": True, "c": BLU}])
notes(s, "0:00-0:15 — Titolo. Una frase: questo strumento non spiega il documento, "
         "lo porta a termine. Tre passi, nessuna registrazione.")

# ─────────────────────────────────────────────── 2. Il problema
s = slide()
kicker(s, Inches(1.0), Inches(0.85), "Il problema")
text(s, Inches(1.0), Inches(1.25), Inches(11.3), Inches(1.0),
     [{"t": "Il modulo non è difficile. È scritto in una lingua che non è la tua.",
       "sz": 34, "b": True}])

b = box(s, Inches(1.0), Inches(2.6), Inches(11.3), Inches(1.5), AMBRA_BG, AMBRA)
text(s, Inches(1.45), Inches(2.9), Inches(10.4), Inches(1.0),
     [{"t": "«Estremi del titolo di occupazione dell'alloggio ai sensi dell'art. 5 D.L. 47/2014»",
       "sz": 21, "i": True, "c": AMBRA, "ls": 1.2}])
text(s, Inches(1.45), Inches(3.55), Inches(10.4), Inches(0.4),
     [{"t": "Campo reale di una dichiarazione di residenza comunale", "sz": 12, "c": GRIGIO}])

for i, (t1, t2) in enumerate([
    ("Si ferma prima di iniziare",
     "Non sa quale informazione le stia chiedendo, non come scriverla."),
    ("Chiede aiuto a qualcuno",
     "Un figlio, un patronato, un CAF. Ogni volta perde autonomia."),
    ("Il modulo resta bianco",
     "E con lui la pratica: la residenza, il contributo, il diritto."),
]):
    x = Inches(1.0 + i * 3.85)
    rule(s, x, Inches(4.55), Inches(3.4), BLU, 2)
    text(s, x, Inches(4.8), Inches(3.4), Inches(0.45),
         [{"t": t1, "sz": 18, "b": True}])
    text(s, x, Inches(5.35), Inches(3.4), Inches(1.1),
         [{"t": t2, "sz": 14, "c": GRIGIO, "ls": 1.3}])
notes(s, "0:15-0:45 — Il problema non è la persona, è il linguaggio. Leggere ad alta "
         "voce l'etichetta reale: nessuno in sala sa cosa chiede. È un campo vero di un "
         "modulo comunale. Il risultato è sempre lo stesso: o chiede aiuto, o rinuncia.")

# ─────────────────────────────────────────────── 3. La soluzione
s = slide(BIANCO)
kicker(s, Inches(1.0), Inches(0.8), "La soluzione")
text(s, Inches(1.0), Inches(1.2), Inches(7.0), Inches(1.5),
     [{"t": "Una domanda alla volta,\nin italiano semplice.", "sz": 34, "b": True, "ls": 1.15}])
for i, (n, t1, t2) in enumerate([
    ("1", "Carica il PDF", "Un solo pulsante. Nessuna registrazione, nessuna configurazione."),
    ("2", "Rispondi", "Una domanda per schermata, in linguaggio comune."),
    ("3", "Scarica il modulo", "Il PDF compilato, pronto da stampare o inviare."),
]):
    y = Inches(2.75 + i * 1.35)
    c = box(s, Inches(1.0), y, Inches(0.62), Inches(0.62), BLU, radius=False)
    text(s, Inches(1.0), y + Inches(0.12), Inches(0.62), Inches(0.4),
         [{"t": n, "sz": 20, "b": True, "c": BIANCO}], align=PP_ALIGN.CENTER)
    text(s, Inches(1.85), y + Inches(0.02), Inches(5.6), Inches(0.4),
         [{"t": t1, "sz": 20, "b": True}])
    text(s, Inches(1.85), y + Inches(0.5), Inches(5.6), Inches(0.6),
         [{"t": t2, "sz": 13.5, "c": GRIGIO, "ls": 1.25}])
pic(s, f"{BUILD}/upload-crop.png", Inches(7.9), Inches(2.1), Inches(4.6), Inches(3.6))
notes(s, "0:45-1:15 — Tre passi. L'interfaccia è un pulsante: niente account, niente "
         "setup. Poi una domanda per schermata, testo grande, un solo campo. "
         "Alla fine il documento pronto.")

# ─────────────────────────────────────────────── 4. Prima / dopo
s = slide()
kicker(s, Inches(1.0), Inches(0.75), "Il cuore del prodotto")
text(s, Inches(1.0), Inches(1.12), Inches(11.3), Inches(0.7),
     [{"t": "Semplificare senza tradire il significato", "sz": 34, "b": True}])

b = box(s, Inches(1.0), Inches(2.1), Inches(5.3), Inches(1.75), BIANCO,
        RGBColor(0xD8, 0xD8, 0xD4))
text(s, Inches(1.35), Inches(2.35), Inches(4.7), Inches(0.3),
     [{"t": "TESTO DEL MODULO", "sz": 11, "b": True, "c": GRIGIO}])
text(s, Inches(1.35), Inches(2.75), Inches(4.7), Inches(0.9),
     [{"t": "Indirizzo di nuova dimora abituale — Via/Piazza", "sz": 17, "ls": 1.2}])

ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(6.55), Inches(2.75),
                        Inches(0.55), Inches(0.4))
ar.fill.solid(); ar.fill.fore_color.rgb = BLU
ar.line.fill.background(); ar.shadow.inherit = False

b = box(s, Inches(7.35), Inches(2.1), Inches(4.95), Inches(1.75), BLU)
text(s, Inches(7.7), Inches(2.35), Inches(4.3), Inches(0.3),
     [{"t": "LA DOMANDA CHE LEGGE LA PERSONA", "sz": 11, "b": True, "c": BLU_CH}])
text(s, Inches(7.7), Inches(2.75), Inches(4.3), Inches(0.9),
     [{"t": "Qual è l'indirizzo della tua nuova casa?", "sz": 19, "b": True, "c": BIANCO, "ls": 1.2}])

text(s, Inches(1.0), Inches(4.25), Inches(5.3), Inches(1.9),
     [{"t": "Il testo originale resta sempre a schermo.", "sz": 17, "b": True, "sa": 8},
      {"t": "La domanda semplice non sostituisce il modulo: lo affianca. "
            "Su un documento con valore legale, una semplificazione che altera "
            "il senso produce una dichiarazione sbagliata — e la responsabilità "
            "resta di chi firma.", "sz": 13.5, "c": GRIGIO, "ls": 1.35}])
pic(s, f"{SHOTS}/03-domanda-indirizzo.png", Inches(6.9), Inches(4.1), Inches(5.4), Inches(2.75))
notes(s, "1:15-1:50 — Questa è la slide che conta. A sinistra il campo com'è scritto "
         "nel modulo, a destra quello che legge la persona. Il testo originale non "
         "sparisce mai: resta sotto, sempre visibile. Non stiamo riscrivendo il "
         "documento, stiamo facendo da ponte.")

# ─────────────────────────────────────────────── 5. Il risultato
s = slide(BIANCO)
kicker(s, Inches(1.0), Inches(0.8), "Il risultato")
text(s, Inches(1.0), Inches(1.18), Inches(6.2), Inches(1.4),
     [{"t": "Non una spiegazione.\nIl documento finito.", "sz": 34, "b": True, "ls": 1.15}])
text(s, Inches(1.0), Inches(2.75), Inches(5.4), Inches(2.2),
     [{"t": "Le risposte finiscono nei campi giusti del PDF originale, che viene "
            "restituito compilato e pronto.", "sz": 15, "c": GRIGIO, "ls": 1.35, "sa": 14},
      {"t": "Prima: il modulo resta bianco.\nDopo: il modulo è compilato.",
       "sz": 16, "b": True, "c": VERDE, "ls": 1.3}])
text(s, Inches(1.0), Inches(5.55), Inches(5.4), Inches(0.8),
     [{"t": "Il guadagno è misurabile senza interpretazioni: o la pratica è "
            "completa, o non lo è.", "sz": 13, "i": True, "c": GRIGIO, "ls": 1.3}])
pic(s, f"{BUILD}/pdf-compilato-crop.png", Inches(6.8), Inches(1.5), Inches(5.6), Inches(4.6))
notes(s, "1:50-2:15 — Il valore non è nella chat, è qui. L'utente non riceve un "
         "riassunto né una guida: riceve il modulo compilato. Prima bianco, dopo "
         "completo. È l'unica metrica che conta e non ha bisogno di interpretazione.")

# ─────────────────────────────────────────────── 6. Lo stack
s = slide()
kicker(s, Inches(1.0), Inches(0.75), "Come è fatto")
text(s, Inches(1.0), Inches(1.12), Inches(11.3), Inches(0.7),
     [{"t": "Tre pezzi, nessuna dipendenza esterna", "sz": 34, "b": True}])

cards = [
    ("FRONTEND", "Angular 21", "Due schermate: carica e rispondi.\nTesto grande, un campo per "
     "volta, navigabile da tastiera.", BLU),
    ("BACKEND", "Spring Boot · Java 21", "Legge i campi del PDF, orchestra le domande, "
     "riscrive il modulo compilato con Apache PDFBox.", BLU),
    ("MOTORE LINGUISTICO", "Ollama · qwen2.5 in locale", "Trasforma l'etichetta burocratica "
     "in una domanda semplice. Gira sulla macchina, non nel cloud.", VERDE),
]
for i, (k, titolo, desc, col) in enumerate(cards):
    x = Inches(1.0 + i * 3.85)
    box(s, x, Inches(2.15), Inches(3.4), Inches(2.55), BIANCO, RGBColor(0xDD, 0xDD, 0xD9))
    rule(s, x, Inches(2.15), Inches(3.4), col, 4)
    text(s, x + Inches(0.3), Inches(2.5), Inches(2.8), Inches(0.3),
         [{"t": k, "sz": 10.5, "b": True, "c": col}])
    text(s, x + Inches(0.3), Inches(2.85), Inches(2.8), Inches(0.5),
         [{"t": titolo, "sz": 18, "b": True}])
    text(s, x + Inches(0.3), Inches(3.5), Inches(2.8), Inches(1.1),
         [{"t": desc, "sz": 13, "c": GRIGIO, "ls": 1.3}])

box(s, Inches(1.0), Inches(5.15), Inches(11.3), Inches(1.35), RGBColor(0xEC, 0xF3, 0xFA))
text(s, Inches(1.4), Inches(5.4), Inches(10.5), Inches(0.35),
     [{"t": "L'INTEGRAZIONE, IN UNA RIGA", "sz": 10.5, "b": True, "c": BLU}])
text(s, Inches(1.4), Inches(5.78), Inches(10.5), Inches(0.5),
     [{"t": "PDF caricato  →  campi estratti  →  ogni etichetta diventa una domanda  →  "
            "risposte riscritte nel PDF  →  download", "sz": 15, "b": True, "ls": 1.2}])
notes(s, "2:15-2:40 — Stack volutamente ordinario e noioso: Angular davanti, Spring Boot "
         "dietro, PDFBox per il PDF. L'unico pezzo interessante è il motore linguistico, "
         "ed è locale. Il flusso è lineare: nessuna coda, nessun database, nessuno stato "
         "da gestire oltre la sessione.")

# ─────────────────────────────────────────────── 7. Perché conta
s = slide(BIANCO)
kicker(s, Inches(1.0), Inches(0.8), "Perché conta")
text(s, Inches(1.0), Inches(1.2), Inches(11.3), Inches(0.8),
     [{"t": "I dati non escono mai dalla macchina", "sz": 34, "b": True}])

punti = [
    ("Nessun account, per nessuno",
     "Né per chi usa lo strumento, né per chi lo installa. Nessun abbonamento AI."),
    ("Il modello gira in locale",
     "Nome, indirizzo, codice fiscale non lasciano il computer. Non è una promessa "
     "contrattuale: è architettura."),
    ("Funziona su qualsiasi modulo compilabile",
     "Non è legato a un modulo specifico: legge i campi di qualunque PDF con moduli."),
    ("Se il modello non risponde, si continua lo stesso",
     "Lo strumento mostra l'etichetta originale e la persona arriva comunque in fondo."),
]
for i, (t1, t2) in enumerate(punti):
    y = Inches(2.35 + i * 1.05)
    d = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.0), y + Inches(0.12),
                           Inches(0.16), Inches(0.16))
    d.fill.solid(); d.fill.fore_color.rgb = BLU
    d.line.fill.background(); d.shadow.inherit = False
    text(s, Inches(1.45), y, Inches(10.8), Inches(0.35),
         [{"t": t1, "sz": 18, "b": True}])
    text(s, Inches(1.45), y + Inches(0.4), Inches(10.6), Inches(0.5),
         [{"t": t2, "sz": 13.5, "c": GRIGIO, "ls": 1.25}])

rule(s, Inches(1.0), Inches(6.45), Inches(11.3), BLU, 2)
text(s, Inches(1.0), Inches(6.7), Inches(11.3), Inches(0.5),
     [{"t": "Un campo alla volta — dal modulo bianco alla pratica conclusa.",
       "sz": 17, "b": True, "c": BLU}])
notes(s, "2:40-3:00 — Chiusura sul punto che vende davvero: privacy per costruzione. "
         "I dati anagrafici non lasciano il computer perché il modello è locale, non "
         "perché qualcuno lo promette in un contratto. Nessun account, nessun "
         "abbonamento, e se il modello cade la persona arriva in fondo comunque. "
         "Chiudere qui.")

prs.save(OUT)
print(f"OK  {OUT}")
print(f"slide: {len(prs.slides._sldIdLst)}")
